"""
core.py — منطق تحليل تقرير "حالة العميل" (PDF/ASPX) وتحويله لإكسيل، وبناء صفحة مساعد التصفية.
نفس الكود المستخدم في النسخة المكتبية (client_status_suite.py)، من غير أي حاجة خاصة بـ tkinter
أو فتح متصفح على جهاز المستخدم — عشان يشتغل على سيرفر (Streamlit) وعلى الجهاز الشخصي مع بعض.
"""
import json
import re
from datetime import date
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ─────────────────────────── أدوات النص العربي ───────────────────────────
AR_RUN = re.compile(r"[\u0600-\u06FF]+")


_RUNS = re.compile(r"[\u0600-\u06FF]+|[0-9A-Za-z_]+(?:[:/.,][0-9A-Za-z_]+)*|.", re.S)


def to_logical(word: str) -> str:
    """الـ PDF بيخزّن الكلمات العربية بالمقلوب (ترتيب بصري) - نرجّعها للترتيب الطبيعي."""
    if not AR_RUN.search(word):
        return word
    runs = _RUNS.findall(word)
    return "".join(r[::-1] if AR_RUN.fullmatch(r) else r for r in reversed(runs))


_NORM = str.maketrans({"أ": "", "إ": "", "آ": "", "ا": "", "ى": "ي", "ة": "ه"})


def norm(s: str) -> str:
    """تطبيع للمقارنة: نشيل الألف بأنواعها لأن Crystal بيفقدها بعد اللام (لا/لأ/لإ)."""
    s = re.sub(r"[\u064B-\u065F\u0640]", "", s)
    return s.translate(_NORM)


# إصلاحات لقيم معروفة فقدت الألف في الـ PDF
VALUE_FIXES = [("ل يوجد", "لا يوجد"), ("ليوجد", "لا يوجد"), ("الهلى", "الأهلى")]


def fix_value(v: str) -> str:
    for a, b in VALUE_FIXES:
        v = v.replace(a, b)
    return v.strip()


# ─────────────────────────── تعريف الحقول ───────────────────────────
CLIENT_LABELS = [
    "الكود", "تاريخ الإصدار", "الاسم", "الرقم القومى", "نوع الاقتراض", "تاريخ التقدم",
    "الفرع", "المندوب الحالى", "موقع التعامل", "رقم الحساب", "رقم حساب الفرع", "الحالة",
    "حالة التعامل مع العميل", "عدد القروض", "عدد القروض المفتوحة", "الحد الائتماني",
    "المبلغ المستخدم", "المبلغ المتاح", "اجمالى الفوائد المستحقة", "اجمالى الغرامات المستحقة",
    "الغرامات المسددة", "تاريخ الميلاد", "النوع", "الحالة الاجتماعية", "المؤهل الدراسي",
    "تليفون محمول", "تليفون المنزل", "المحافظة", "المركز", "القرية", "عنوان المنزل",
    "تصنيف الإقامة", "اسم المنشأة", "البطاقة الضريبية", "السجل التجاري", "السجل الصناعي",
    "تليفون العمل", "منطقة العمل", "قطاع العمل", "نوع النشاط", "مجال التخصص", "عنوان العمل",
    "البريد الإلكترونى", "الموقع الإلكترونى", "ملاحظات",
]
LOAN_LABELS = [
    "نوع القرض", "طريقة الحساب", "كود القرض", "العملة", "حالة القرض", "تاريخ القرض",
    "قيمة القرض", "عدد الأقساط", "فترة السماح", "أيام الترحيل", "المصاريف الموزعة",
    "المصاريف المقدمة", "رسوم طلب القرض", "عمولة المندوب", "مصاريف القسط", "رسوم الطوابع",
    "فترة السداد كل", "مندوب التنمية", "الحالة القانونية", "غرامات مستحقة", "غرامات مسددة",
    "غرامات معفاة", "رصيد العميل", "أيام التأخير", "أيام التبكير",
]
GUAR_LABELS = [
    "كود الضامن", "الاسم", "الرقم القومى", "النوع", "تاريخ الإصدار", "تاريخ الميلاد",
    "الموبايل", "الرقم البريدى", "العنوان",
]
INST_LABELS = ["رصيد العميل", "أيام التأخير", "أيام التبكير"]
IGNORE_LABELS = ["خطوط الطول", "خطوط العرض", "قيمة مستقلة لا تستقطع من المصاريف الموزعة"]


def _compile(labels, ignore=()):
    """قائمة (كلمات مطبّعة، الاسم الأصلي، هل نتجاهله) مرتبة من الأطول للأقصر."""
    items = [(tuple(norm(w) for w in l.split()), l, False) for l in labels]
    items += [(tuple(norm(w) for w in l.split()), l, True) for l in ignore]
    return sorted(items, key=lambda x: -len(x[0]))


SETS = {
    "client": _compile(CLIENT_LABELS, IGNORE_LABELS),
    "loan": _compile(LOAN_LABELS, IGNORE_LABELS),
    "guarantor": _compile(GUAR_LABELS),
    "installments": _compile(INST_LABELS),
    "loans_list": [],
}

INST_COLS = [
    "رقم القسط", "تاريخ الاستحقاق", "أصل القسط", "مصاريف القسط", "إجمالي القسط",
    "أصل المسدد", "مصاريف المسدد", "إجمالي المسدد", "الحالة", "تاريخ الحالة",
    "بنك / خزينة التحصيل", "تأخير / تبكير (أيام)",
]
LOANLIST_SCHEMAS = {
    10: ["رقم القرض", "التاريخ", "القيمة", "عدد الأقساط", "فترة السماح", "فترة السداد",
         "نسبة المصاريف الموزعة %", "نسبة المصاريف المقدمة %", "إجمالي المصاريف", "حالة القرض"],
    13: ["رقم القرض", "التاريخ", "القيمة", "عدد الأقساط", "أقساط مجدولة", "فترة السماح",
         "فترة السداد", "نسبة المصاريف الموزعة %", "نسبة المصاريف المقدمة %",
         "إجمالي المصاريف", "حالة القرض", "عدد أقساط متأخرة", "أيام تأخير"],
}

DATE_RE = re.compile(r"^(\d{4})/(\d{2})/(\d{2})$")
NUMLIKE = re.compile(r"^[\d/.,%:\-]+$")


def to_date(s):
    m = DATE_RE.match(str(s))
    if not m:
        return s
    try:
        return date(int(m[1]), int(m[2]), int(m[3]))
    except ValueError:
        return s


def to_num(s):
    t = str(s).replace(",", "").replace("%", "").strip()
    if re.fullmatch(r"-?\d+", t):
        return int(t)
    if re.fullmatch(r"-?\d+\.\d+", t):
        return float(t)
    return s


# ─────────────────────────── قراءة الـ PDF ───────────────────────────
def page_lines(page):
    """يرجّع سطور الصفحة، وكل سطر قائمة خلايا (كل خلية قائمة كلمات بالترتيب المنطقي، يمين لشمال)."""
    words = page.extract_words(x_tolerance=2, y_tolerance=2)
    for w in words:
        w["t"] = to_logical(w["text"])
        w["cy"] = (w["top"] + w["bottom"]) / 2
    words.sort(key=lambda w: w["cy"])
    lines = []
    for w in words:
        if lines and abs(w["cy"] - lines[-1][0]) <= 3:
            lines[-1][1].append(w)
        else:
            lines.append([w["cy"], [w]])
    out = []
    for cy, ws in lines:
        ws.sort(key=lambda w: -w["x1"])
        cells, prev = [], None
        for w in ws:
            if prev is not None and prev["x0"] - w["x1"] <= 9:
                cells[-1].append(w["t"])
            else:
                cells.append([w["t"]])
            prev = w
        out.append((cy, cells))
    return out


# حقول قيمتها بتمتد على أكتر من خلية (رقم + وحدته)
MULTI_CELL = {"المصاريف الموزعة", "المصاريف المقدمة", "فترة السداد كل"}


def match_label(words, i, label_set):
    """هل الكلمات من i تبدأ بعنوان حقل معروف؟"""
    nw = tuple(norm(w) for w in words[i:i + 8])
    for lw, name, ignored in label_set:
        if nw[:len(lw)] == lw:
            return name, len(lw), ignored
    return None


def assign_fields(cells, label_set):
    """
    يطلّع (اسم الحقل، القيمة) من خلايا السطر.
    الـ PDF عربي (يمين->شمال) فالقيمة بتيجي بعد العنوان: في نفس الخلية أو في الخلية اللي بعدها.
    """
    fields, current, value_cell = [], None, None
    for ci, cell in enumerate(cells):
        i = 0
        while i < len(cell):
            # العنوان يبدأ من أول الخلية، أو بعد رقم/تاريخ في نفس الخلية (زي: "تاريخ الإصدار 2024/11/18 الكود 1/9431")
            allow = i == 0 or bool(NUMLIKE.match(cell[i - 1]))
            m = match_label(cell, i, label_set) if allow else None
            if m:
                name, n, ignored = m
                current, value_cell = None, None
                if not ignored:
                    current = [name, []]
                    fields.append(current)
                i += n
                continue
            if current is not None:
                if not current[1]:
                    value_cell = ci
                if value_cell == ci or current[0] in MULTI_CELL:
                    current[1].append(cell[i])
            i += 1
    return [(n, fix_value(" ".join(v))) for n, v in fields]


def check_is_pdf(file):
    """
    الملف .aspx اللي بيتحفظ من المتصفح غالباً PDF جواه. نتأكد قبل القراءة عشان نطلّع رسالة مفهومة.
    file: مسار على القرص، أو أي كائن بيه read()/seek() (زي الملف المرفوع في Streamlit).
    """
    if hasattr(file, "read"):
        pos = file.tell() if hasattr(file, "tell") else None
        first = file.read(1024)
        if hasattr(file, "seek"):
            file.seek(pos if pos is not None else 0)
    else:
        with open(file, "rb") as f:
            first = f.read(1024)
    if b"%PDF-" not in first:
        raise ValueError(
            "الملف ده مش تقرير PDF (غالباً اتحفظ كصفحة ويب HTML). "
            "افتح التقرير في المتصفح واحفظه بـ Ctrl+S أو من زرار الحفظ اللي في عارض الـ PDF.")


MONEY_RE = re.compile(r"^-?[\d,]+\.\d{2}$")
INT_RE = re.compile(r"^-?\d+$")


def parse_installment_row(cells):
    """
    صف قسط: م، تاريخ الاستحقاق، (أصل، مصاريف، إجمالي) القسط، (أصل، مصاريف، إجمالي) المسدد،
    الحالة (كلمة أو اتنين)، تاريخ الحالة، بنك/خزينة التحصيل، أيام التأخير/التبكير.
    بنقرأه كلمة كلمة عشان الخلايا القريبة من بعض ما تبوّظش الترتيب.
    """
    w = [x for c in cells for x in c]
    if len(w) < 12 or not w[0].isdigit() or not DATE_RE.match(w[1]):
        return None
    money = w[2:8]
    if not all(MONEY_RE.match(x) for x in money) or not INT_RE.match(w[-1]):
        return None
    mid = w[8:-1]
    di = next((i for i, x in enumerate(mid) if DATE_RE.match(x)), None)
    if di is None:
        status, sdate, bank = " ".join(mid), "", ""
    else:
        status, sdate, bank = " ".join(mid[:di]), mid[di], fix_value(" ".join(mid[di + 1:]))
    return [w[0], w[1]] + money + [fix_value(status), sdate, bank, w[-1]]


def parse_totals_row(cells):
    """صف الإجمالي: 'إجمالي' + 6 أرقام، وبعدها (رصيد العميل / أيام التأخير / أيام التبكير)."""
    toks = [(ci, x) for ci, c in enumerate(cells) for x in c]
    if not toks or norm(toks[0][1]) != norm("إجمالي"):
        return None
    nums = [x for _, x in toks[1:7]]
    if len(nums) < 6 or not all(MONEY_RE.match(x) for x in nums):
        return None
    rest = {}
    for ci, x in toks[7:]:
        rest.setdefault(ci, []).append(x)
    return nums, [rest[k] for k in sorted(rest)]


def parse_pdf(path):
    """
    path: مسار ملف على القرص، أو كائن ملف في الذاكرة (bytes/BytesIO/الملف المرفوع في Streamlit).
    """
    check_is_pdf(path)
    if hasattr(path, "seek"):
        path.seek(0)
    client, loans, guarantors, installments, loan_list, unknown = {}, [], [], [], [], []
    section, list_title = "client", ""
    carry = None  # عنوان بدون قيمة في آخر السطر - قيمته في أول السطر اللي بعده
    with pdfplumber.open(path) as pdf:
        for pno, page in enumerate(pdf.pages, start=1):
            carry, prev_cells = None, None      # (القسم بيكمّل من الصفحة اللي فاتت: الجدول ممكن يتكمل في صفحة تانية)
            for cy, cells in page_lines(page):
                # هيدر الصفحة: في الصفحة الأولى فيه بيانات العميل الفعلية (من y=65)، وفي الباقي بيتكرر (نتخطاه لحد y=90)
                if cy < (65 if pno == 1 else 90):
                    continue
                flat = " ".join(" ".join(c) for c in cells)
                if "رقم الصفحة" in flat or cy > 835:      # فوتر الصفحة
                    continue
                nflat = norm(flat)

                # عناوين الأقسام
                if "بيانات" in flat and norm("القروض المفتوحة") in nflat:
                    section = "loan"; continue
                if norm("بيانات الضامنون") in nflat:
                    section = "guarantor"; continue
                if norm("بيانات القرض التفصيلية") in nflat:
                    section = "installments"; continue
                if nflat.startswith("قروض") or nflat.startswith(norm("ملخص القروض")):
                    section, list_title = "loans_list", flat.strip(); continue

                # خلايا "N شهر" و "رقم تاريخ" المدموجة
                fixed = []
                for c in cells:
                    txt = " ".join(c)
                    m = re.match(r"^(\d+) (\d{4}/\d{2}/\d{2})$", txt)
                    if m:
                        fixed += [[m[1]], [m[2]]]
                    else:
                        fixed.append(c)
                cells = fixed
                texts = [" ".join(c) for c in cells]

                # جدول الأقساط
                if section == "installments":
                    loan_code = loans[-1].get("كود القرض", "") if loans else ""
                    row = parse_installment_row(cells)
                    if row:
                        installments.append([loan_code] + row)
                        continue
                    tot = parse_totals_row(cells)
                    if tot:
                        nums, rest_cells = tot
                        installments.append([loan_code, "إجمالي", ""] + nums + ["", "", "", ""])
                        for n, v in assign_fields(rest_cells, SETS["installments"]):
                            if loans:
                                loans[-1][n] = v
                        continue

                # جدول القروض (غير مصدرة / مسددة)
                if section == "loans_list":
                    merged = []
                    for t in texts:
                        if t == "شهر" and merged and merged[-1].isdigit():
                            merged[-1] += " شهر"
                        else:
                            merged.append(fix_value(t))
                    if len(merged) >= 10 and merged[0].isdigit() and DATE_RE.match(merged[1]):
                        cols = LOANLIST_SCHEMAS.get(len(merged))
                        if cols is None:
                            cols = [f"عمود {k}" for k in range(1, len(merged) + 1)]
                            unknown.append((pno, merged))
                        loan_list.append(dict(zip(["القسم"] + cols, [list_title] + merged)))
                    continue

                if section not in SETS or not SETS[section]:
                    continue

                # حالة خاصة: "القروض الملغاة / المسددة" عنوانهم متقسم على سطرين
                if [norm(t) for t in texts] == [norm("الملغاة"), norm("المسددة")] and prev_cells and len(prev_cells) >= 2:
                    client.setdefault("القروض الملغاة", " ".join(prev_cells[-2]))
                    client.setdefault("القروض المسددة", " ".join(prev_cells[-1]))
                    prev_cells = cells
                    continue

                # قيمة معلّقة من السطر اللي فات (زي: اجمالى الفوائد المستحقة)
                if carry and cells:
                    first = " ".join(cells[0])
                    if not match_label(cells[0], 0, SETS[section]):
                        target, label = carry
                        if not target.get(label):
                            target[label] = fix_value(first)
                        cells = cells[1:]
                    carry = None
                prev_cells = list(cells)

                fields = assign_fields(cells, SETS[section])
                if not fields:
                    continue

                if section == "client":
                    for n, v in fields:
                        client.setdefault(n, v)
                    if fields[-1][1] == "":
                        carry = (client, fields[-1][0])
                elif section == "loan":
                    for n, v in fields:
                        if n == "نوع القرض" and (not loans or "نوع القرض" in loans[-1]):
                            loans.append({})
                        if not loans:
                            loans.append({})
                        loans[-1][n] = v
                elif section == "guarantor":
                    for n, v in fields:
                        if n == "كود الضامن":
                            guarantors.append({"صفحة": pno})
                        if not guarantors:
                            guarantors.append({"صفحة": pno})
                        guarantors[-1][n] = v
                elif section == "installments":
                    for n, v in fields:
                        if loans:
                            loans[-1][n] = v
    return client, loans, guarantors, installments, loan_list, unknown


# ─────────────────────────── كتابة الإكسيل ───────────────────────────
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFF")
NOTE = ("ملاحظة: النص العربي متاخد من ملف PDF، وCrystal Reports بيفقد الألف بعد اللام في بعض الكلمات "
        "(لا / لأ / لإ). القيم المعروفة اتصلحت، لكن راجع الأسماء والعناوين اللي فيها الحرف لا.")


def _style_sheet(ws, widths=None):
    ws.sheet_view.rightToLeft = True
    for col in ws.columns:
        L = get_column_letter(col[0].column)
        longest = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[L].width = min(max(10, longest + 3), 48)
        for c in col:
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="right")


def _table(ws, headers, rows, date_cols=(), num_cols=()):
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    for r in rows:
        r = list(r)
        for i in date_cols:
            if i < len(r):
                r[i] = to_date(r[i])
        for i in num_cols:
            if i < len(r) and r[i] != "":
                r[i] = to_num(r[i])
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for i in date_cols:
            if i < len(row) and isinstance(row[i].value, date):
                row[i].number_format = "yyyy/mm/dd"
    ws.freeze_panes = "A2"


def write_workbook(result, out_path):
    client, loans, guarantors, installments, loan_list, unknown = result
    wb = Workbook()

    ws = wb.active
    ws.title = "بيانات العميل"
    _table(ws, ["البيان", "القيمة"], [[k, v] for k, v in client.items()])
    ws.append([])
    ws.append([NOTE])
    _style_sheet(ws)

    if loans:
        ws = wb.create_sheet("بيانات القرض")
        keys = []
        for l in loans:
            keys += [k for k in l if k not in keys]
        _table(ws, keys, [[l.get(k, "") for k in keys] for l in loans])
        _style_sheet(ws)

    if loan_list:
        ws = wb.create_sheet("ملخص القروض")
        keys = []
        for l in loan_list:
            keys += [k for k in l if k not in keys]
        rows = [[l.get(k, "") for k in keys] for l in loan_list]
        dcols = [i for i, k in enumerate(keys) if k == "التاريخ"]
        ncols = [i for i, k in enumerate(keys) if k not in ("القسم", "التاريخ", "فترة السداد", "حالة القرض") and not k.startswith("عمود")]
        _table(ws, keys, rows, date_cols=dcols, num_cols=ncols)
        _style_sheet(ws)

    if guarantors:
        ws = wb.create_sheet("الضامنون")
        keys = []
        for g in guarantors:
            keys += [k for k in g if k not in keys]
        _table(ws, keys, [[g.get(k, "") for k in keys] for g in guarantors],
               date_cols=[i for i, k in enumerate(keys) if k in ("تاريخ الإصدار", "تاريخ الميلاد")])
        _style_sheet(ws)

    if installments:
        ws = wb.create_sheet("الأقساط")
        _table(ws, ["كود القرض"] + INST_COLS, installments,
               date_cols=[2, 10], num_cols=[1, 3, 4, 5, 6, 7, 8, 12])
        _style_sheet(ws)

    if unknown:
        ws = wb.create_sheet("غير معروف")
        for pno, row in unknown:
            ws.append([pno] + row)

    wb.save(out_path)


def convert(pdf_path, xlsx_path=None):
    pdf_path = Path(pdf_path)
    xlsx_path = Path(xlsx_path) if xlsx_path else pdf_path.with_suffix(".xlsx")
    result = parse_pdf(pdf_path)
    write_workbook(result, xlsx_path)
    return xlsx_path, result


# ─────────────────────────── بيانات مساعد التصفية ───────────────────────────
def installments_by_loan(installments):
    """يقسّم صفوف الأقساط حسب كود القرض (من غير صف الإجمالي)."""
    groups = {}
    for r in installments:
        if r[1] == "إجمالي":
            continue
        groups.setdefault(r[0], []).append(r)
    return groups


def liquidation_lines(rows):
    """
    يحوّل الأقساط لنفس شكل السطور اللي كانت بتتنسخ من برنامج المحصل،
    فمساعد التصفية يقراها بنفس دالة التحليل بتاعته من غير أي تعديل فيه.
    """
    out = []
    for r in rows:
        _code, m, due, a, b, c, pa, pb, pc, status, sdate, bank, delay = r
        parts = [m, due, a, b, c, pa, pb, pc, status or "-", sdate, bank, delay]
        out.append(" ".join(str(x) for x in parts if str(x).strip() != ""))
    return "\n".join(out)


def office_from_branch(branch):
    b = (branch or "").strip()
    return b[3:].strip() if b.startswith("فرع ") else b


def build_liquidation_html(template_html, rows, client, source_name="", loan_code=""):
    """
    يرجّع نص صفحة مساعد التصفية بعد ما يحط فيها بيانات القرض تلقائي،
    بإضافة سكربت صغير في آخرها من غير ما يلمس كود الصفحة الأصلي.
    """
    payload = {
        "lines": liquidation_lines(rows),
        "clientName": client.get("الاسم", ""),
        "clientCode": client.get("الكود", ""),
        "officeName": office_from_branch(client.get("الفرع", "")),
        "source": source_name,
        "loan": loan_code,
    }
    data = json.dumps(payload, ensure_ascii=True).replace("</", "<\\/")
    script = """
<script>
/* ===== دمج البرنامج الجديد: تحميل الأقساط من ملف التقرير بدل اللصق اليدوي ===== */
(function () {
  var P = %s;
  function set(id, v) { var e = document.getElementById(id); if (e && v) e.value = v; }
  var box = document.getElementById('dataInput');
  if (!box) return;
  var info = document.createElement('div');
  info.style.cssText = 'background:#eafaf1;border:1px solid #1e8449;color:#145a32;border-radius:8px;padding:8px 12px;margin-bottom:10px;font-size:13px;text-align:right;';
  info.textContent = '\\u2705 الأقساط اتحمّلت من ملف التقرير' + (P.source ? ' (' + P.source + ')' : '') +
    (P.clientName ? ' — العميل: ' + P.clientName : '') + (P.clientCode ? ' — كود ' + P.clientCode : '') +
    (P.loan ? ' — القرض ' + P.loan : '') + '. تقدر تعدّل النص تحت وتعمل تحليل تاني لو حبيت.';
  box.parentNode.insertBefore(info, box);
  box.value = P.lines;
  try { processAll(); } catch (e) { console.error(e); }
  set('clientName', P.clientName); set('clientCode', P.clientCode); set('officeName', P.officeName);
})();
</script>
""" % data
    idx = template_html.rfind("</body>")
    if idx == -1:
        return template_html + script
    return template_html[:idx] + script + template_html[idx:]
